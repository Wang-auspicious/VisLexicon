"""Quiet Keynote：一套克制的 16:9 演示模板（封面 / 章节 / 大数字 / 左右对比 / 图文 / 结尾）。

用法（Agent 可原样执行）:
    python build_deck.py                                  # 生成 quiet-keynote.pptx（示例内容）
    python build_deck.py --content deck.json --out my.pptx  # 用自己的内容，格式见 sample-content.json
    python build_deck.py --render                         # 额外用本机 PowerPoint 把每页导出成 PNG（仅 Windows + Office）

只依赖 python-pptx。所有文字都是真文本框，可在 PowerPoint / Keynote / WPS 里继续改。
"""
import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
W, H = Emu(12192000), Emu(6858000)  # 13.333 × 7.5 in，16:9
INK, SUB, BG, CARD, ACCENT, LINE = "1D1D1F", "6E6E73", "F5F5F7", "FFFFFF", "0071E3", "D2D2D7"
LATIN, CJK = "Segoe UI", "Microsoft YaHei UI"


def rgb(h): return RGBColor.from_string(h)


def text(slide, x, y, w, h, s, size, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing: p.line_spacing = spacing
        r = p.add_run(); r.text = line
        f = r.font; f.size = Pt(size); f.bold = bold; f.color.rgb = rgb(color); f.name = LATIN
        rPr = r._r.get_or_add_rPr()
        ea = rPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}ea", {"typeface": CJK})
        rPr.append(ea)
    return box


def bg(slide, color=BG):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = rgb(color)


def rect(slide, x, y, w, h, color, radius=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(color); shp.line.fill.background(); shp.shadow.inherit = False
    if radius: shp.adjustments[0] = radius
    return shp


def footer(slide, deck, n):
    text(slide, Emu(762000), H - Emu(520000), Emu(6000000), Emu(260000), deck.get("footer", ""), 10, SUB)
    text(slide, W - Emu(1762000), H - Emu(520000), Emu(1000000), Emu(260000), f"{n:02d}", 10, SUB, align=PP_ALIGN.RIGHT)


M = Emu(762000)  # 左右边距 0.83in


def s_cover(prs, d, deck, n):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    text(s, M, Emu(2100000), W - 2 * M, Emu(1500000), d["title"], 54, bold=True, spacing=1.0)
    text(s, M, Emu(3700000), W - 2 * M, Emu(700000), d.get("subtitle", ""), 22, SUB)
    rect(s, M, Emu(4700000), Emu(560000), Emu(56000), ACCENT)
    text(s, M, Emu(4900000), W - 2 * M, Emu(400000), d.get("meta", ""), 12, SUB)


def s_section(prs, d, deck, n):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, INK)
    text(s, M, Emu(2500000), Emu(2000000), Emu(500000), d.get("kicker", ""), 14, "86868B")
    text(s, M, Emu(2950000), W - 2 * M, Emu(1300000), d["title"], 44, "F5F5F7", bold=True)


def s_number(prs, d, deck, n):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    text(s, M, Emu(900000), W - 2 * M, Emu(500000), d.get("kicker", ""), 16, SUB)
    text(s, M, Emu(1500000), W - 2 * M, Emu(2600000), d["number"], 150, ACCENT, bold=True, spacing=0.9)
    text(s, M, Emu(4300000), Emu(7600000), Emu(1200000), d.get("caption", ""), 22, INK, spacing=1.25)
    footer(s, deck, n)


def s_compare(prs, d, deck, n):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    text(s, M, Emu(700000), W - 2 * M, Emu(800000), d["title"], 32, bold=True)
    gap = Emu(300000); cw = int((W - 2 * M - gap) / 2); top = Emu(1800000); ch = Emu(4000000)
    for i, col in enumerate(d["columns"][:2]):
        x = M + i * (cw + gap)
        rect(s, x, top, cw, ch, CARD, radius=0.06)
        text(s, x + Emu(420000), top + Emu(380000), cw - Emu(840000), Emu(500000), col["label"], 14, ACCENT if i else SUB, bold=True)
        text(s, x + Emu(420000), top + Emu(950000), cw - Emu(840000), Emu(2800000), "\n".join(col["points"]), 18, INK, spacing=1.45)
    footer(s, deck, n)


def s_image(prs, d, deck, n):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    text(s, M, Emu(700000), W - 2 * M, Emu(800000), d["title"], 32, bold=True)
    img = (HERE / d["image"]).resolve()
    card = rect(s, M, Emu(1700000), W - 2 * M, Emu(3900000), CARD, radius=0.03)
    pic = s.shapes.add_picture(str(img), M + Emu(300000), Emu(1950000), height=Emu(3400000))
    maxw = W - 2 * M - Emu(600000)
    if pic.width > maxw:  # 等比缩到卡片里
        ratio = maxw / pic.width; pic.width = int(maxw); pic.height = int(pic.height * ratio)
    pic.left = int(M + (W - 2 * M - pic.width) / 2); pic.top = int(Emu(1700000) + (Emu(3900000) - pic.height) / 2)
    text(s, M, Emu(5750000), W - 2 * M, Emu(400000), d.get("caption", ""), 12, SUB)
    footer(s, deck, n)


def s_end(prs, d, deck, n):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    text(s, M, Emu(2600000), W - 2 * M, Emu(1200000), d["title"], 48, bold=True, align=PP_ALIGN.CENTER)
    text(s, M, Emu(3800000), W - 2 * M, Emu(500000), d.get("subtitle", ""), 18, SUB, align=PP_ALIGN.CENTER)


LAYOUTS = {"cover": s_cover, "section": s_section, "number": s_number, "compare": s_compare, "image": s_image, "end": s_end}


def build(content, out):
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    for i, sl in enumerate(content["slides"], 1):
        LAYOUTS[sl["layout"]](prs, sl, content, i)
    prs.save(out); print("wrote", out, len(content["slides"]), "slides")


def render(pptx, outdir):
    """用本机 PowerPoint 导出每一页 PNG（真实渲染，不是示意）。"""
    import win32com.client
    outdir.mkdir(exist_ok=True)
    app = win32com.client.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(str(pptx.resolve()), ReadOnly=True, WithWindow=False)
    try:
        for i, slide in enumerate(pres.Slides, 1):
            slide.Export(str((outdir / f"slide-{i:02d}.png").resolve()), "PNG", 1600, 900)
    finally:
        pres.Close()
    print("rendered", outdir)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--content", default=str(HERE / "sample-content.json"))
    ap.add_argument("--out", default=str(HERE / "quiet-keynote.pptx"))
    ap.add_argument("--render", action="store_true")
    a = ap.parse_args()
    out = Path(a.out)
    build(json.loads(Path(a.content).read_text(encoding="utf-8")), out)
    if a.render: render(out, out.parent / "previews")
